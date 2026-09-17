#!/usr/bin/env python3
"""Genera los .docx de las plantillas del cliente a partir de los .md.

Uso:
    python3 _tools/plantillas_docx.py            # todas
    python3 _tools/plantillas_docx.py turnos     # una

Fuente: _plantillas/*.md   Salida: _plantillas/docx/plantilla-<nombre>.docx
Editar siempre el .md; el .docx se regenera.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ACCENT = RGBColor(0xC8, 0x55, 0x2B)
INK = RGBColor(0x1E, 0x1B, 0x18)
MUTED = RGBColor(0x6B, 0x64, 0x5D)
LINE = "E8E3DD"
BAND = "FAF8F5"

SALIDA = {
    "qr-links": "plantilla-qr-links.docx",
    "menu-digital": "plantilla-menu-digital.docx",
    "catalogo": "plantilla-catalogo.docx",
    "mini-web": "plantilla-mini-web.docx",
    "perfil-google": "plantilla-perfil-google.docx",
    "turnos": "plantilla-turnos.docx",
}


def _shade(cell, color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(el)


def _borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), LINE)
        borders.append(el)
    tbl_pr.append(borders)


def _inline(par, texto):
    """Escribe texto con **negrita**, `código` y casillas."""
    texto = texto.replace("[ ]", "☐")
    for trozo in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            run = par.add_run(trozo[2:-2])
            run.bold = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            run = par.add_run(trozo[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            run = par.add_run(trozo)
        run.font.color.rgb = INK


def _fila(linea):
    celdas = [c.strip() for c in linea.strip().strip("|").split("|")]
    return celdas


def construir(md_path, docx_path):
    lineas = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(1.8)

    i = 0
    while i < len(lineas):
        linea = lineas[i]

        if linea.startswith("|") and i + 1 < len(lineas) and set(lineas[i + 1].replace("|", "").strip()) <= set("-: "):
            encabezado = _fila(linea)
            i += 2
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                filas.append(_fila(lineas[i]))
                i += 1
            ancho = max([len(encabezado)] + [len(f) for f in filas])
            tabla = doc.add_table(rows=1, cols=ancho)
            tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
            tabla.autofit = True
            _borders(tabla)
            for j, txt in enumerate(encabezado):
                cel = tabla.rows[0].cells[j]
                cel.text = ""
                _shade(cel, BAND)
                par = cel.paragraphs[0]
                _inline(par, txt)
                for run in par.runs:
                    run.bold = True
                    run.font.size = Pt(10)
            for fila in filas:
                celdas = tabla.add_row().cells
                for j in range(ancho):
                    txt = fila[j] if j < len(fila) else ""
                    celdas[j].text = ""
                    par = celdas[j].paragraphs[0]
                    if txt:
                        _inline(par, txt)
                    for run in par.runs:
                        run.font.size = Pt(10)
            doc.add_paragraph()
            continue

        if linea.startswith("# "):
            par = doc.add_paragraph()
            run = par.add_run(linea[2:].strip())
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = ACCENT
            par.space_after = Pt(6)
        elif linea.startswith("## "):
            par = doc.add_paragraph()
            par.space_before = Pt(14)
            run = par.add_run(linea[3:].strip().upper())
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = ACCENT
        elif linea.strip().startswith("- "):
            par = doc.add_paragraph(style="List Bullet")
            _inline(par, linea.strip()[2:])
            for run in par.runs:
                run.font.size = Pt(10.5)
        elif linea.strip():
            par = doc.add_paragraph()
            _inline(par, linea.strip())
            if linea.strip().startswith("Los ") or linea.strip().startswith("Todo "):
                for run in par.runs:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = MUTED
        i += 1

    pie = doc.add_paragraph()
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pie.add_run("Gianluca Salvatori · Servicios digitales · WhatsApp +54 11 3784-5392")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    base = Path(__file__).resolve().parent.parent / "_plantillas"
    if not base.exists():
        base = Path.cwd()
    pedidas = sys.argv[1:] or list(SALIDA)
    for nombre in pedidas:
        md = base / f"{nombre}.md"
        if not md.exists():
            sys.exit(f"No existe {md}")
        print(construir(md, base / "docx" / SALIDA[nombre]))
